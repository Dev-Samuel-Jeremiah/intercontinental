"""
results/views_tools.py
──────────────────────
Grading scale, positions and score-entry deadlines; scores by Excel; the
class broadsheet with subject analysis.
"""
import io
from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify

from academics.models import Term
from academics.utils import get_active_session, get_active_term
from users.models import ENROLLED_STATUSES, Class, Student, Subject
from users.views_profile import can_manage, can_view

from . import analysis, locks
from .assessment import AssessmentError, clean_score
from .models import AssessmentScheme, ResultSheetSettings, TermResult

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _deny(request, who):
    messages.error(request, f'Only {who}.')
    return redirect('dashboard:router')


def _can_enter(user):
    from .views import _can_enter_scores
    return _can_enter_scores(user)


# ─────────────────────────────────────────────────────────────────────────────
# Grading, positions, deadlines
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def grading(request):
    if not can_manage(request.user):
        return _deny(request, 'administrators can change grading and deadlines')
    rs, _ = ResultSheetSettings.objects.get_or_create(tenant=request.tenant)
    active_session = get_active_session()
    terms = list(Term.objects.filter(tenant=request.tenant, session=active_session).order_by('start_date', 'name')) \
        if active_session else []
    if request.method == 'POST':
        action = request.POST.get('action')
        try:
            if action in ('bands', 'preset'):
                rows = (analysis.PRESETS[request.POST['preset']][1] if action == 'preset'
                        else analysis.parse_bands(request.POST))
                with transaction.atomic():
                    analysis.save_bands(request.tenant, rows)
                    changed = analysis.regrade(request.tenant)
                messages.success(request, f'Grading scale saved. {changed} saved result{"s" if changed != 1 else ""} '
                                          'regraded.' if changed else 'Grading scale saved.')
            elif action == 'positions':
                rs.show_position = bool(request.POST.get('show_position'))
                rs.show_subject_position = bool(request.POST.get('show_subject_position'))
                rs.save(update_fields=['show_position', 'show_subject_position'])
                messages.success(request, 'Report card positions updated.')
            elif action == 'deadline':
                term = get_object_or_404(Term, pk=request.POST.get('term'))
                raw = request.POST.get('locks_on')
                if raw:
                    locks.set_deadline(term, date.fromisoformat(raw))
                    messages.success(request, f'Score entry for {term.name} closes after {raw}.')
                else:
                    locks.lock_for(term) and locks.lock_for(term).delete()
                    messages.success(request, f'No deadline for {term.name}.')
            elif action == 'reopen':
                lock = locks.lock_for(get_object_or_404(Term, pk=request.POST.get('term')))
                locks.reopen(lock, request.user, request.POST.get('reason'), request.POST.get('hours') or 48)
                messages.success(request, f'Score entry reopened until {timezone.localtime(lock.reopened_until):%d %b %-I:%M %p}.')
            elif action == 'close':
                locks.close_now(locks.lock_for(get_object_or_404(Term, pk=request.POST.get('term'))))
                messages.success(request, 'Score entry closed again.')
        except (analysis.GradingError, locks.ScoresLocked, ValueError, KeyError) as problem:
            messages.error(request, str(problem) or 'Check the values and try again.')
        return redirect('results:grading')
    band_list = analysis.bands(request.tenant)
    rows = [{'grade': b.grade, 'min': b.min_score, 'desc': b.description} for b in band_list]
    rows += [{'grade': '', 'min': '', 'desc': ''}] * max(3, 12 - len(rows))
    return render(request, 'results/grading.html', {
        'rows': rows, 'presets': analysis.PRESETS, 'has_bands': bool(band_list), 'settings': rs,
        'terms': [{'term': t, 'lock': locks.lock_for(t)} for t in terms], 'session': active_session,
    })


# ─────────────────────────────────────────────────────────────────────────────
# Scores by Excel
# ─────────────────────────────────────────────────────────────────────────────

def _sheet_target(request, data):
    school_class = get_object_or_404(Class, pk=data.get('class'))
    subject = get_object_or_404(Subject, pk=data.get('subject'))
    term = get_object_or_404(Term.objects.select_related('session'), pk=data.get('term'))
    return school_class, subject, term


@login_required
def score_sheet(request):
    """A pre-filled Excel sheet for one class, subject and term."""
    if not _can_enter(request.user):
        return _deny(request, 'staff allowed to enter scores')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from curriculum.progression import offers
    school_class, subject, term = _sheet_target(request, request.GET)
    scheme = AssessmentScheme.for_tenant(request.tenant)
    cols = scheme.columns
    existing = {tr.student_id: tr for tr in TermResult.objects.filter(
        tenant=request.tenant, class_assigned=school_class, subject=subject, term=term, session=term.session)}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Scores'
    ws.append([f'{request.tenant.name} — {subject.name}, {school_class.name}, {term.session.name} {term.name}'])
    ws.append([f'Fill the score columns, then upload this file. Do not change the first two columns or rows. '
               f'Put Y under "Not offering" for a student who does not take {subject.name}.'])
    header = ['Admission no.', 'Student'] + [f"{c['label']} ({c['max']})" for c in cols] + ['Not offering (Y)']
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='7C5C1E')
    ws['A1'].font = Font(bold=True, size=13)
    students = Student.objects.filter(tenant=request.tenant, class_assigned=school_class,
                                      status__in=ENROLLED_STATUSES).order_by('full_name')
    for st in students:
        tr = existing.get(st.id)
        not_off = tr.is_not_offering if tr else not offers(st, subject.id)
        ws.append([st.admission_number, st.full_name]
                  + [(getattr(tr, c['field']) or None) if tr and not not_off else None for c in cols]
                  + ['Y' if not_off else ''])
    last = ws.max_row
    for i, c in enumerate(cols):
        letter = ws.cell(row=3, column=3 + i).column_letter
        dv = DataValidation(type='decimal', operator='between', formula1='0', formula2=str(c['max']),
                            allow_blank=True, showErrorMessage=True,
                            error=f"{c['label']} is out of {c['max']}.", errorTitle='Score too high')
        ws.add_data_validation(dv)
        dv.add(f'{letter}4:{letter}{max(last, 4)}')
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 30
    for i in range(len(cols) + 1):
        ws.column_dimensions[ws.cell(row=3, column=3 + i).column_letter].width = 16
    ws.freeze_panes = 'C4'
    buf = io.BytesIO()
    wb.save(buf)
    name = slugify(f'{school_class.name}-{subject.name}-{term.name}') or 'scores'
    resp = HttpResponse(buf.getvalue(), content_type=XLSX)
    resp['Content-Disposition'] = f'attachment; filename="{name}.xlsx"'
    return resp


def read_sheet(upload, school_class, subject, term, scheme):
    """[(student, {field: score}, not_offering)] or AssessmentError listing every problem."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(upload, data_only=True, read_only=True)
    except Exception:
        raise AssessmentError("That isn't an Excel file this page can read — download the sheet and fill that in.")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    cols = scheme.columns
    header_at = next((i for i, r in enumerate(rows[:6]) if r and str(r[0] or '').strip().lower().startswith('admission')), None)
    if header_at is None:
        raise AssessmentError('The sheet has no "Admission no." header — use the downloaded sheet.')
    header = [str(h or '').strip() for h in rows[header_at]]
    expected = [f"{c['label']} ({c['max']})" for c in cols]
    if header[2:2 + len(cols)] != expected:
        raise AssessmentError('The columns don\'t match the current assessment structure. Download a fresh sheet.')
    roster = {s.admission_number.strip().lower(): s for s in Student.objects.filter(
        tenant_id=school_class.tenant_id, class_assigned=school_class, status__in=ENROLLED_STATUSES)}
    out, problems, seen = [], [], set()
    for n, r in enumerate(rows[header_at + 1:], start=header_at + 2):
        if not r or not any(v not in (None, '') for v in r):
            continue
        adm = str(r[0] or '').strip()
        st = roster.get(adm.lower())
        if not st:
            problems.append(f'Row {n}: "{adm}" is not a student in {school_class.name}.')
            continue
        if st.id in seen:
            problems.append(f'Row {n}: {st.full_name} appears twice.')
            continue
        seen.add(st.id)
        flag = str(r[2 + len(cols)] if len(r) > 2 + len(cols) else '').strip().upper()
        not_off = flag in ('Y', 'YES', 'X')
        scores = {}
        if not not_off:
            for i, c in enumerate(cols):
                raw = r[2 + i] if len(r) > 2 + i else None
                try:
                    scores[c['field']] = clean_score('' if raw is None else str(raw).strip(), c['max'], c['label'])
                except AssessmentError as problem:
                    problems.append(f'Row {n} ({st.full_name}): {problem}')
        out.append((st, scores, not_off))
    if problems:
        raise AssessmentError(' '.join(problems[:8]) + (f' …and {len(problems) - 8} more.' if len(problems) > 8 else ''))
    if not out:
        raise AssessmentError('The sheet has no students in it.')
    return out


@transaction.atomic
def save_sheet(rows, school_class, subject, term, scheme):
    saved = 0
    for st, scores, not_off in rows:
        tr, _ = TermResult.objects.get_or_create(tenant_id=school_class.tenant_id, student=st, subject=subject,
                                                 term=term, session=term.session,
                                                 defaults={'class_assigned': school_class})
        tr.class_assigned = school_class
        tr.is_not_offering = not_off
        for col in AssessmentScheme.COLUMNS:
            setattr(tr, f'{col}_score', scores.get(f'{col}_score', 0) if not not_off else 0)
        tr.save()
        saved += 1
    return saved


@login_required
def score_upload(request):
    if not _can_enter(request.user):
        return _deny(request, 'staff allowed to enter scores')
    school_class, subject, term = _sheet_target(request, request.POST)
    back = f"{reverse('results:manual_score_entry')}?class={school_class.id}&subject={subject.id}&term={term.id}"
    if request.method != 'POST' or not request.FILES.get('sheet'):
        messages.error(request, 'Choose the filled-in Excel sheet to upload.')
        return redirect(back)
    try:
        locks.check_open(term)
        from .release import may_edit_scores
        allowed, why = may_edit_scores(request.user, school_class, subject, term)
        if not allowed:
            raise AssessmentError(why)
        scheme = AssessmentScheme.for_tenant(request.tenant)
        rows = read_sheet(request.FILES['sheet'], school_class, subject, term, scheme)
        n = save_sheet(rows, school_class, subject, term, scheme)
        messages.success(request, f'Scores uploaded for {n} student{"s" if n != 1 else ""}.')
    except (AssessmentError, locks.ScoresLocked) as problem:
        messages.error(request, f'Nothing was saved. {problem}')
    return redirect(back)


# ─────────────────────────────────────────────────────────────────────────────
# Broadsheet and subject analysis
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def broadsheet(request):
    from .views import _is_admin_or_super
    staff = getattr(request.user, 'staff', None)
    see_all = _is_admin_or_super(request.user) or getattr(staff, 'can_view_all_results', False)
    classes = Class.objects.filter(tenant=request.tenant).order_by('name')
    if not see_all:
        classes = classes.filter(form_teachers=staff) if staff else classes.none()
    if not can_view(request.user) or not classes.exists():
        return _deny(request, 'administrators, staff allowed to view all results, and form teachers (for their '
                              'own class) can see broadsheets')
    terms = Term.objects.filter(tenant=request.tenant).select_related('session').order_by('-start_date')
    school_class = classes.filter(pk=request.GET.get('class')).first() if request.GET.get('class') else classes.first()
    term = terms.filter(pk=request.GET.get('term')).first() if request.GET.get('term') else (get_active_term() or terms.first())
    data = analysis.broadsheet(school_class, term, term.session) if school_class and term else None
    if data and request.GET.get('format') == 'xlsx':
        return _broadsheet_xlsx(request, school_class, term, data)
    return render(request, 'results/broadsheet.html', {
        'classes': classes, 'terms': terms, 'school_class': school_class, 'term': term, 'data': data,
    })


def _broadsheet_xlsx(request, school_class, term, data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Broadsheet'
    ws.append([f'{request.tenant.name} — {school_class.name} broadsheet, {term.session.name} {term.name}'])
    ws['A1'].font = Font(bold=True, size=13)
    head = ['Pos.', 'Admission no.', 'Student'] + [s.name for s in data['subjects']] + ['Subjects', 'Total', 'Average']
    ws.append(head)
    for cell in ws[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='7C5C1E')
        cell.alignment = Alignment(text_rotation=90 if cell.column > 3 and cell.column <= 3 + len(data['subjects']) else 0,
                                   horizontal='center')
    for r in data['rows']:
        ws.append([analysis.ordinal(r['position']), r['student'].admission_number, r['student'].full_name]
                  + [c['score'] if c else None for c in r['cells']]
                  + [r['count'], r['total'], r['average']])
    ws.append([])
    ws.append(['', '', 'Subject mean'] + [next((a['mean'] for a in data['analysis'] if a['subject'] == s), None)
                                          for s in data['subjects']])
    ws.append(['', '', 'Pass rate %'] + [next((a['pass_rate'] for a in data['analysis'] if a['subject'] == s), None)
                                         for s in data['subjects']])
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.freeze_panes = 'D3'
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type=XLSX)
    resp['Content-Disposition'] = f'attachment; filename="{slugify(school_class.name + "-" + term.name)}-broadsheet.xlsx"'
    return resp
