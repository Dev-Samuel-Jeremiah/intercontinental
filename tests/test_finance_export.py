"""
The accounting export: balanced journals in the school's own account codes,
as an Excel workbook and as CSV for QuickBooks, Xero and the rest.
"""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.test import override_settings
from django.utils import timezone

from finance.models import (AccountingCodes, Expense, ExpenseCategory, Payment,
                            PettyCashTopUp, Refund, Wallet)
from finance.services import export, refunds as refund_service
from tenants.middleware import set_current_tenant
from users.models import User

from .factories import TenantTestCase, make_admin, make_school, make_student

START, END = date(2026, 9, 1), date(2026, 9, 30)


def _aware(day):
    return timezone.make_aware(datetime.combine(day, datetime.min.time().replace(hour=10)))


class ExportBase(TenantTestCase):

    def setUp(self):
        super().setUp()
        self.school = make_school('alpha')
        set_current_tenant(self.school)
        self.student = make_student(self.school, 'STU001', full_name='Ade Bello')
        self.wallet = Wallet.objects.create(tenant=self.school, student=self.student)
        self.diesel = ExpenseCategory.objects.create(tenant=self.school, name='Diesel')

    def _payment(self, amount, method='BANK_TRANSFER', day=date(2026, 9, 10),
                 status=Payment.APPROVED, reference=None):
        return Payment.objects.create(
            tenant=self.school, wallet=self.wallet, amount=Decimal(str(amount)),
            method=method, status=status, created_at=_aware(day),
            reference=reference or f'PAY-{Payment.objects.count() + 1:04d}')

    def _expense(self, amount, method=Expense.TRANSFER, status=Expense.PAID,
                 day=date(2026, 9, 12), category=None):
        return Expense.objects.create(tenant=self.school, category=category or self.diesel,
                                      description='Generator diesel', payee='Jos Diesel',
                                      amount=Decimal(str(amount)), method=method,
                                      status=status, spent_on=day)

    def _paid_refund(self, amount, method=Refund.TRANSFER, day=date(2026, 9, 20)):
        refund = refund_service.request_refund(self.school, self.student, amount, method=method)
        refund_service.approve(refund)
        refund_service.mark_paid(refund, paid_on=day)
        return refund

    def _lines(self, entry):
        return {(l['code'], 'Dr' if l['debit'] else 'Cr'): l['debit'] or l['credit']
                for l in entry['lines']}


class JournalTests(ExportBase):

    def test_a_bank_transfer_payment(self):
        self._payment(50000)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(self._lines(entry), {('1000', 'Dr'): Decimal('50000'),
                                              ('4000', 'Cr'): Decimal('50000')})
        self.assertIn('Ade Bello', entry['description'])

    def test_each_payment_method_lands_in_its_own_account(self):
        for method in ('CASH', 'PAYSTACK', 'POS', 'CHEQUE'):
            self._payment(1000, method=method)
        debits = sorted(e['lines'][0]['code'] for e in export.journal(self.school, START, END))
        self.assertEqual(debits, ['1000', '1000', '1010', '1020'])

    def test_pending_and_rejected_payments_are_left_out(self):
        self._payment(1000, status=Payment.PENDING)
        self._payment(1000, status=Payment.REJECTED)
        self.assertEqual(export.journal(self.school, START, END), [])

    def test_only_the_period_is_exported(self):
        self._payment(1000, day=date(2026, 8, 31))
        self._payment(2000, day=date(2026, 9, 1))
        self._payment(3000, day=date(2026, 9, 30))
        self._payment(4000, day=date(2026, 10, 1))
        amounts = [e['lines'][0]['debit'] for e in export.journal(self.school, START, END)]
        self.assertEqual(amounts, [Decimal('2000'), Decimal('3000')])

    def test_a_paid_expense(self):
        self._expense(185000)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(self._lines(entry), {('6000', 'Dr'): Decimal('185000'),
                                              ('1000', 'Cr'): Decimal('185000')})

    def test_a_category_with_its_own_code_uses_it(self):
        self.diesel.account_code = '6120'
        self.diesel.save()
        self._expense(185000)
        line = export.journal(self.school, START, END)[0]['lines'][0]
        self.assertEqual((line['code'], line['account']), ('6120', 'Diesel'))

    def test_petty_cash_spending_comes_out_of_petty_cash(self):
        self._expense(3500, method=Expense.PETTY_CASH)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(entry['lines'][1]['code'], '1030')

    def test_unpaid_expenses_are_left_out(self):
        self._expense(5000, status=Expense.PENDING)
        self._expense(5000, status=Expense.APPROVED)
        self._expense(5000, status=Expense.REJECTED)
        self.assertEqual(export.journal(self.school, START, END), [])

    def test_a_petty_cash_top_up_moves_money_from_the_bank(self):
        PettyCashTopUp.objects.create(tenant=self.school, amount=Decimal('20000'),
                                      topped_up_on=date(2026, 9, 5))
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(self._lines(entry), {('1030', 'Dr'): Decimal('20000'),
                                              ('1000', 'Cr'): Decimal('20000')})

    def test_a_refund_by_transfer(self):
        self._paid_refund(12000)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(self._lines(entry), {('4900', 'Dr'): Decimal('12000'),
                                              ('1000', 'Cr'): Decimal('12000')})

    def test_a_refund_to_the_wallet_is_owed_to_the_parent_not_cash_out(self):
        self._paid_refund(8000, method=Refund.WALLET)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(entry['lines'][1]['code'], '2100')

    def test_every_entry_balances(self):
        self._payment(50000); self._payment(12500.50, method='CASH')
        self._expense(185000); self._expense(3500, method=Expense.PETTY_CASH)
        self._paid_refund(12000)
        PettyCashTopUp.objects.create(tenant=self.school, amount=Decimal('20000'),
                                      topped_up_on=date(2026, 9, 5))
        entries = export.journal(self.school, START, END)
        self.assertEqual(len(entries), 6)
        for entry in entries:
            self.assertEqual(sum(l['debit'] for l in entry['lines']),
                             sum(l['credit'] for l in entry['lines']))
        debit, credit = export.totals(entries)
        self.assertEqual(debit, credit)

    def test_entries_are_dated_in_order_and_numbered(self):
        self._expense(1000, day=date(2026, 9, 20))
        self._payment(2000, day=date(2026, 9, 3))
        entries = export.journal(self.school, START, END)
        self.assertEqual([e['date'] for e in entries], [date(2026, 9, 3), date(2026, 9, 20)])
        self.assertEqual([e['number'] for e in entries], ['TC2609-0001', 'TC2609-0002'])

    def test_the_schools_own_codes_are_used(self):
        codes = AccountingCodes.for_school(self.school)
        codes.bank, codes.bank_name = '1100', 'Zenith fees account'
        codes.fee_income = '4100'
        codes.save()
        self._payment(50000)
        entry = export.journal(self.school, START, END)[0]
        self.assertEqual(entry['lines'][0]['account'], 'Zenith fees account')
        self.assertEqual(self._lines(entry), {('1100', 'Dr'): Decimal('50000'),
                                              ('4100', 'Cr'): Decimal('50000')})

    def test_the_trial_balance_nets_each_account(self):
        self._payment(50000)
        self._expense(20000)
        rows = {r['code']: r for r in export.trial_balance(export.journal(self.school, START, END))}
        self.assertEqual(rows['1000']['debit'] - rows['1000']['credit'], Decimal('30000'))
        self.assertEqual(rows['4000']['credit'], Decimal('50000'))
        self.assertEqual(rows['6000']['debit'], Decimal('20000'))

    def test_another_schools_money_is_never_exported(self):
        self._payment(50000)
        other = make_school('beta')
        set_current_tenant(other)
        theirs = make_student(other, 'STU001')
        wallet = Wallet.objects.create(tenant=other, student=theirs)
        Payment.objects.create(tenant=other, wallet=wallet, amount=Decimal('999999'),
                               reference='PAY-BETA', method='CASH', status=Payment.APPROVED,
                               created_at=_aware(date(2026, 9, 10)))
        set_current_tenant(None)
        entries = export.journal(self.school, START, END)
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]['lines'][0]['debit'], Decimal('50000'))


class FileTests(ExportBase):

    def setUp(self):
        super().setUp()
        self._payment(50000, reference='PAY-AAA1')
        self._expense(185000)

    def _rows(self, fmt):
        return list(csv.reader(io.StringIO(export.journal_csv(self.school, START, END, fmt))))

    def test_the_generic_journal_csv(self):
        rows = self._rows('generic')
        self.assertEqual(rows[0], ['Date', 'Journal No', 'Reference', 'Account Code',
                                   'Account Name', 'Description', 'Debit', 'Credit'])
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[1][:5], ['2026-09-10', 'TC2609-0001', 'PAY-AAA1', '1000', 'Bank'])
        self.assertEqual(rows[1][6:], ['50000.00', '0.00'])

    def test_the_xero_csv_uses_signed_amounts(self):
        rows = self._rows('xero')
        self.assertEqual(rows[0][:6], ['*Narration', '*Date', 'Description', '*AccountCode',
                                       '*TaxRate', '*Amount'])
        amounts = [Decimal(r[5]) for r in rows[1:]]
        self.assertEqual(sum(amounts), Decimal('0'))
        self.assertEqual(rows[1][1], '10/09/2026')

    def test_the_quickbooks_csv_names_accounts(self):
        rows = self._rows('quickbooks')
        self.assertEqual(rows[0][:5], ['Journal No', 'Journal Date', 'Account', 'Debits',
                                       'Credits'])
        self.assertEqual(rows[1][2], 'Bank')
        self.assertEqual((rows[1][3], rows[1][4]), ('50000.00', ''))
        self.assertEqual((rows[2][3], rows[2][4]), ('', '50000.00'))

    def test_the_workbook_has_every_sheet(self):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(export.workbook(self.school, START, END)))
        self.assertEqual(book.sheetnames, ['Summary', 'Fees received', 'Shop and canteen',
                                           'Refunds', 'Expenses',
                                           'Withholding tax', 'Petty cash top-ups', 'Journal',
                                           'Trial balance'])
        self.assertEqual(book['Fees received']['F2'].value, 50000)
        self.assertEqual(book['Journal'].max_row, 5)
        labels = {row[0].value: row[1].value for row in book['Summary'].iter_rows(min_row=5)
                  if row[0].value}
        self.assertEqual(labels['Surplus / (deficit)'], -135000)
        self.assertEqual(labels['Total debits'], labels['Total credits'])


@override_settings(ROOT_DOMAIN='localhost', ALLOWED_HOSTS=['*'])
class ExportPageTests(ExportBase):

    def setUp(self):
        super().setUp()
        self.admin = make_admin(self.school)
        set_current_tenant(self.school)
        self.client.force_login(self.admin)
        self._payment(50000)

    def _get(self, path):
        return self.client.get(path, HTTP_HOST=self.host(self.school))

    def test_the_page_shows_the_period(self):
        response = self._get('/finance/admin/books/?start=2026-09-01&end=2026-09-30')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['entry_count'], 1)
        self.assertTrue(response.context['balanced'])

    def test_a_reversed_range_is_put_the_right_way_round(self):
        response = self._get('/finance/admin/books/?start=2026-09-30&end=2026-09-01')
        self.assertEqual(response.context['start'], date(2026, 9, 1))

    def test_a_bad_date_falls_back_instead_of_crashing(self):
        self.assertEqual(self._get('/finance/admin/books/?start=nonsense').status_code, 200)

    def test_downloading_the_csv(self):
        response = self._get('/finance/admin/books/journal.csv'
                             '?start=2026-09-01&end=2026-09-30&format=xero')
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('alpha-academy-books-20260901-20260930-journal-xero.csv',
                      response['Content-Disposition'])
        self.assertIn(b'*Narration', response.content)

    def test_downloading_the_workbook(self):
        response = self._get('/finance/admin/books/books.xlsx?start=2026-09-01&end=2026-09-30')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'PK'))

    def test_saving_account_codes(self):
        self.client.post('/finance/admin/books/',
                         {'bank': '1100', 'bank_name': 'Zenith fees account',
                          f'category_{self.diesel.id}': '6120'},
                         HTTP_HOST=self.host(self.school))
        codes = AccountingCodes.objects.get(tenant=self.school)
        self.diesel.refresh_from_db()
        self.assertEqual((codes.bank, codes.bank_name), ('1100', 'Zenith fees account'))
        self.assertEqual(codes.cash, '1010')    # untouched fields keep their default
        self.assertEqual(self.diesel.account_code, '6120')

    def test_a_parent_cannot_download_the_books(self):
        parent = User.objects.create_user(username='par1', password='x',
                                          tenant=self.school, is_parent=True)
        self.client.force_login(parent)
        for path in ('/finance/admin/books/', '/finance/admin/books/journal.csv',
                     '/finance/admin/books/books.xlsx'):
            self.assertNotEqual(self._get(path).status_code, 200, path)
