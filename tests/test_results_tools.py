"""
Grading scales, positions with ties, score-entry deadlines, scores by Excel,
and the class broadsheet with subject analysis.
"""
import io
from datetime import timedelta

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from academics.models import AcademicSession, Term
from results import analysis, locks
from results.assessment import AssessmentError
from results.models import AssessmentScheme, Grading, ResultSheetSettings, ScoreLock, TermResult
from tenants.middleware import set_current_tenant
from users.models import ClassSubject, Staff, Subject

from .factories import TenantTestCase, make_admin, make_class, make_school, make_student, make_teacher


class Base(TenantTestCase):

    def setUp(self):
        super().setUp()
        self.school = make_school('alpha')
        self.admin = make_admin(self.school)
        self.jss1 = make_class(self.school, 'JSS 1A')
        set_current_tenant(self.school)
        today = timezone.localdate()
        self.session = AcademicSession.objects.create(tenant=self.school, name='2026/2027',
                                                      start_date=today - timedelta(days=60), is_active=True)
        self.term = Term.objects.create(tenant=self.school, name='1st Term', session=self.session,
                                        start_date=today - timedelta(days=30), end_date=today + timedelta(days=60),
                                        is_active=True)
        self.maths = Subject.objects.create(tenant=self.school, name='Mathematics')
        self.english = Subject.objects.create(tenant=self.school, name='English')
        for s in (self.maths, self.english):
            ClassSubject.objects.create(tenant=self.school, school_class=self.jss1, subject=s)
        self.kids = [make_student(self.school, f'S{i}', full_name=f'Kid {i}', school_class=self.jss1) for i in range(4)]
        set_current_tenant(self.school)
        analysis.save_bands(self.school, analysis.PRESETS['waec'][1])

    def _get(self, path, **kw):
        return self.client.get(path, HTTP_HOST=self.host(self.school), **kw)

    def _post(self, path, data=None, **kw):
        return self.client.post(path, data or {}, HTTP_HOST=self.host(self.school), **kw)

    def _score(self, kid, subject, total, **kw):
        return TermResult.objects.create(tenant=self.school, student=kid, class_assigned=self.jss1, subject=subject,
                                         session=self.session, term=self.term, ca1_score=min(total, 30),
                                         exam_score=max(0, total - 30), **kw)


class GradingTests(Base):

    def test_rounding_is_half_up(self):
        self.assertEqual((analysis.whole(74.5), analysis.whole(75.5), analysis.whole(74.4)), (75, 76, 74))

    def test_band_rules(self):
        with self.assertRaises(analysis.GradingError):
            analysis.check_bands([('A', 70, ''), ('B', 50, '')])        # nothing starts at 0
        with self.assertRaises(analysis.GradingError):
            analysis.check_bands([('A', 70, ''), ('A', 0, '')])
        with self.assertRaises(analysis.GradingError):
            analysis.check_bands([('A', 50, ''), ('B', 50, ''), ('F', 0, '')])
        for _, (_, rows) in analysis.PRESETS.items():
            analysis.check_bands(rows)

    def test_grades_use_only_this_schools_bands(self):
        other = make_school('beta')
        Grading.objects.create(tenant=other, grade='Z', min_score=0)
        tr = self._score(self.kids[0], self.maths, 72)
        set_current_tenant(None)    # e.g. a background job
        tr.save()
        self.assertEqual(tr.grade, 'B2')

    def test_changing_the_scale_regrades_saved_results(self):
        tr = self._score(self.kids[0], self.maths, 72)
        self.assertEqual(tr.grade, 'B2')
        analysis.save_bands(self.school, analysis.PRESETS['letters'][1])
        self.assertEqual(analysis.regrade(self.school), 1)
        tr.refresh_from_db()
        self.assertEqual((tr.grade, tr.remark), ('A', 'Excellent'))

    def test_grading_page(self):
        self.client.force_login(self.admin)
        self.assertContains(self._get('/results/grading/'), 'A1')
        self._post('/results/grading/', {'action': 'bands', 'grade_0': 'a', 'min_0': '60', 'desc_0': 'Top',
                                         'grade_1': 'F', 'min_1': '0'})
        self.assertEqual(list(Grading.objects.values_list('grade', flat=True)), ['A', 'F'])
        self._post('/results/grading/', {'action': 'bands', 'grade_0': 'A', 'min_0': '60'})   # one band only
        self.assertEqual(Grading.objects.count(), 2)
        self._post('/results/grading/', {'action': 'positions', 'show_position': 'on'})
        rs = ResultSheetSettings.objects.get()
        self.assertEqual((rs.show_position, rs.show_subject_position), (True, False))
        teacher = make_teacher(self.school)
        self.client.force_login(teacher)
        self.assertEqual(self._get('/results/grading/').status_code, 302)


class PositionTests(Base):

    def test_ties_share_a_position(self):
        self.assertEqual(analysis.rank({'a': 90, 'b': 80, 'c': 80, 'd': 70}), {'a': 1, 'b': 2, 'c': 2, 'd': 4})
        self.assertEqual([analysis.ordinal(n) for n in (1, 2, 3, 4, 11, 12, 13, 21, 22, 101)],
                         ['1st', '2nd', '3rd', '4th', '11th', '12th', '13th', '21st', '22nd', '101st'])

    def test_class_and_subject_positions_ignore_not_offering(self):
        k = self.kids
        for kid, m, e in ((k[0], 80, 60), (k[1], 70, 70), (k[2], 90, 40)):
            self._score(kid, self.maths, m)
            self._score(kid, self.english, e)
        self._score(k[3], self.maths, 95)
        self._score(k[3], self.english, 0, is_not_offering=True)
        positions, averages = analysis.class_positions(self.jss1, self.term, self.session)
        self.assertEqual(averages[k[3].id], 95)          # not offering English doesn't drag the average
        self.assertEqual([positions[x.id] for x in k], [2, 2, 4, 1])
        subj = analysis.subject_positions(self.jss1, self.term, self.session)
        self.assertNotIn((k[3].id, self.english.id), subj)
        self.assertEqual(subj[(k[1].id, self.english.id)], 1)

    def test_report_card_shows_positions_when_switched_on(self):
        self._score(self.kids[0], self.maths, 80)
        self._score(self.kids[1], self.maths, 70)
        for tr in TermResult.objects.all():
            tr.published = True
            tr.save()
        self.client.force_login(self.kids[1].user)
        url = f'/results/results/{self.session.id}/{self.term.id}/'
        self.assertNotContains(self._get(url), 'of 2</small>')
        rs, _ = ResultSheetSettings.objects.get_or_create(tenant=self.school)
        rs.show_position = rs.show_subject_position = True
        rs.save()
        r = self._get(url)
        self.assertContains(r, '2nd <small')
        self.assertContains(r, 'of 2</small>')


class LockTests(Base):

    def test_deadline_and_reopen(self):
        lock = locks.set_deadline(self.term, timezone.localdate() - timedelta(days=1))
        self.assertTrue(lock.is_locked)
        with self.assertRaises(locks.ScoresLocked):
            locks.check_open(self.term)
        with self.assertRaises(locks.ScoresLocked):
            locks.reopen(lock, self.admin, '')
        locks.reopen(lock, self.admin, 'Late CBT for SS3', hours=2)
        locks.check_open(self.term)
        locks.close_now(lock)
        self.assertTrue(lock.is_locked)
        locks.set_deadline(self.term, timezone.localdate())   # closes at the end of today
        self.assertFalse(locks.lock_for(self.term).is_locked)

    def test_locked_entry_saves_nothing(self):
        locks.set_deadline(self.term, timezone.localdate() - timedelta(days=1))
        self.client.force_login(self.admin)
        url = f'/results/manual-scores/?class={self.jss1.id}&subject={self.maths.id}&term={self.term.id}'
        self.assertContains(self._get(url), 'closed on')
        tr = TermResult.objects.get(student=self.kids[0], subject=self.maths)
        self._post('/results/manual-scores/', {'class': self.jss1.id, 'subject': self.maths.id, 'term': self.term.id,
                                               f'ca1_{tr.id}': 8, f'exam_{tr.id}': 50})
        tr.refresh_from_db()
        self.assertEqual(tr.total_score, 0)


class ExcelTests(Base):

    def _sheet(self):
        self.client.force_login(self.admin)
        r = self._get(f'/results/scores/sheet/?class={self.jss1.id}&subject={self.maths.id}&term={self.term.id}')
        self.assertEqual(r['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(r.content))

    def _upload(self, wb):
        buf = io.BytesIO()
        wb.save(buf)
        return self._post('/results/scores/upload/', {
            'class': self.jss1.id, 'subject': self.maths.id, 'term': self.term.id,
            'sheet': SimpleUploadedFile('s.xlsx', buf.getvalue())})

    def test_round_trip(self):
        wb = self._sheet()
        ws = wb.active
        self.assertEqual([c.value for c in ws[3]][:6], ['Admission no.', 'Student', 'TRT (10)', 'Homework (10)',
                                                         'Class Part. (10)', 'Exam (70)'])
        self.assertEqual(ws.max_row, 3 + 4)
        for row in range(4, 8):
            ws.cell(row=row, column=3, value=8)
            ws.cell(row=row, column=6, value=55.5)
        ws.cell(row=7, column=7, value='Y')            # last kid doesn't take Maths
        self._upload(wb)
        rows = {tr.student.full_name: tr for tr in TermResult.objects.select_related('student')}
        self.assertEqual(rows['Kid 0'].total_score, 63.5)
        self.assertTrue(rows['Kid 3'].is_not_offering)
        self.assertEqual(rows['Kid 3'].total_score, 0)

    def test_bad_rows_save_nothing(self):
        wb = self._sheet()
        ws = wb.active
        ws.cell(row=4, column=6, value=88)             # exam out of 70
        ws.cell(row=5, column=1, value='NOT-HERE')
        r = self._upload(wb)
        self.assertFalse(TermResult.objects.filter(total_score__gt=0).exists())
        msgs = [str(m) for m in r.wsgi_request._messages]
        self.assertTrue(any('out of 70' in m and 'NOT-HERE' in m for m in msgs), msgs)

    def test_sheet_must_match_the_structure(self):
        wb = self._sheet()
        from results.assessment import PRESETS, save
        save(self.school, PRESETS['two_tests'][1], self.term)
        self._upload(wb)
        self.assertFalse(TermResult.objects.filter(total_score__gt=0).exists())

    def test_locked_upload_is_refused(self):
        wb = self._sheet()
        wb.active.cell(row=4, column=3, value=5)
        locks.set_deadline(self.term, timezone.localdate() - timedelta(days=1))
        self._upload(wb)
        self.assertFalse(TermResult.objects.filter(total_score__gt=0).exists())


class BroadsheetTests(Base):

    def setUp(self):
        super().setUp()
        k = self.kids
        for kid, m, e in ((k[0], 80, 60), (k[1], 30, 45), (k[2], 90, 75)):
            self._score(kid, self.maths, m)
            self._score(kid, self.english, e)

    def test_grid_and_analysis(self):
        data = analysis.broadsheet(self.jss1, self.term, self.session)
        self.assertEqual([s.name for s in data['subjects']], ['English', 'Mathematics'])
        self.assertEqual([r['student'].full_name for r in data['rows']], ['Kid 2', 'Kid 0', 'Kid 1', 'Kid 3'])
        self.assertIsNone(data['rows'][-1]['position'])
        maths = next(a for a in data['analysis'] if a['subject'] == self.maths)
        self.assertEqual((maths['mean'], maths['highest'], maths['lowest'], maths['pass_rate']), (66.7, 90, 30, 67))
        self.assertEqual(dict(maths['distribution'])['A1'], 2)

    def test_pages_and_export(self):
        self.client.force_login(self.admin)
        r = self._get(f'/results/broadsheet/?class={self.jss1.id}&term={self.term.id}')
        self.assertContains(r, 'Kid 2')
        self.assertContains(r, 'Subject performance')
        r = self._get(f'/results/broadsheet/?class={self.jss1.id}&term={self.term.id}&format=xlsx')
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['C3'].value, 'Kid 2')
        self.assertEqual(ws['A3'].value, '1st')

    def test_only_form_teachers_see_their_class(self):
        teacher = make_teacher(self.school)
        self.client.force_login(teacher)
        self.assertEqual(self._get('/results/broadsheet/').status_code, 302)
        set_current_tenant(self.school)
        self.jss1.form_teachers.set([Staff.objects.get(user=teacher)])
        self.assertEqual(self._get('/results/broadsheet/').status_code, 200)


class ReportAccessTests(Base):

    def test_a_parent_only_sees_their_own_childs_report(self):
        from users.models import User
        mine, theirs = self.kids[0], self.kids[1]
        theirs.parent_email = 'someone-else@example.com'
        theirs.save()
        for kid in (mine, theirs):
            self._score(kid, self.maths, 70, published=True)
        parent = User.objects.create_user(username=f'parent_{mine.id}', password='x', tenant=self.school,
                                          is_parent=True, email=mine.parent_email)
        self.client.force_login(parent)
        ok = self._get(f'/results/results/{mine.id}/{self.session.id}/{self.term.id}/')
        self.assertEqual(ok.status_code, 200)
        self.assertContains(ok, 'Kid 0')
        blocked = self._get(f'/results/results/{theirs.id}/{self.session.id}/{self.term.id}/')
        self.assertEqual(blocked.status_code, 403)

    def test_unpublished_results_stay_hidden(self):
        self._score(self.kids[0], self.maths, 70)
        self.client.force_login(self.kids[0].user)
        self.assertContains(self._get(f'/results/results/{self.session.id}/{self.term.id}/'), 'not been published')
