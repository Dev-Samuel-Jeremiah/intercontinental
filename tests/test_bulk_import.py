"""
Bulk registration of students and staff from Excel: reading files, checking
every row, importing the good ones, the new logins working, and permissions.
"""
import csv
import io
from datetime import date, datetime

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile

from saas_admin.models import Campus
from users import bulk_import as bi
from users.models import Staff, Student, User

from .factories import TenantTestCase, make_admin, make_class, make_school, make_student, make_teacher

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def xlsx(rows, name='list.xlsx'):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type=XLSX)


class Base(TenantTestCase):

    def setUp(self):
        super().setUp()
        self.school = make_school('alpha')
        self.admin = make_admin(self.school)
        self.jss1 = make_class(self.school, 'JSS 1A')

    def _post(self, path, data=None, user=None):
        self.client.force_login(user or self.admin)
        return self.client.post(path, data or {}, HTTP_HOST=self.host(self.school))

    def _get(self, path, user=None):
        self.client.force_login(user or self.admin)
        return self.client.get(path, HTTP_HOST=self.host(self.school))


class CleaningTests(Base):

    def test_phone_numbers(self):
        self.assertEqual(bi.clean_phone(8031234567.0), '08031234567')          # Excel dropped the 0
        self.assertEqual(bi.clean_phone('+234 803 123 4567'), '08031234567')
        self.assertEqual(bi.clean_phone('0805-555-1234'), '08055551234')
        with self.assertRaises(ValueError):
            bi.clean_phone('12345')

    def test_dates(self):
        self.assertEqual(bi.clean_date('14/03/2014'), date(2014, 3, 14))
        self.assertEqual(bi.clean_date('2014-03-14'), date(2014, 3, 14))
        self.assertEqual(bi.clean_date(datetime(2014, 3, 14)), date(2014, 3, 14))
        self.assertEqual(bi.clean_date('41712'), date(2014, 3, 14))           # arrived as an Excel number
        for bad in ('31/02/2012', '2090-01-01', 'soon'):
            with self.assertRaises(ValueError):
                bi.clean_date(bad)

    def test_headers_are_matched_loosely_and_blank_lines_skipped(self):
        rows, unknown = bi.read_rows(xlsx([['Student Name', 'Adm No', 'Favourite Colour'],
                                           ['Amina Bello', 'A1', 'blue'], [None, None, None]]), 'students')
        self.assertEqual(rows, [(2, {'full_name': 'Amina Bello', 'admission_number': 'A1'})])
        self.assertEqual(unknown, ['Favourite Colour'])

    def test_files_that_cannot_be_used(self):
        with self.assertRaisesMessage(bi.ImportFileError, 'Excel file'):
            bi.read_rows(SimpleUploadedFile('a.docx', b'x'), 'students')
        with self.assertRaisesMessage(bi.ImportFileError, 'Full Name'):
            bi.read_rows(xlsx([['Surname', 'Class'], ['Bello', 'JSS 1A']]), 'students')
        with self.assertRaisesMessage(bi.ImportFileError, 'no names'):
            bi.read_rows(xlsx([['Full Name']]), 'students')

    def test_csv(self):
        buf = io.StringIO()
        csv.writer(buf).writerows([['Full Name', 'Class'], ['Amina Bello', 'JSS 1A']])
        rows, _ = bi.read_rows(SimpleUploadedFile('a.csv', buf.getvalue().encode()), 'students')
        self.assertEqual(rows[0][1]['class_name'], 'JSS 1A')


class StudentTests(Base):

    def check(self, rows, create_classes=False):
        parsed, _ = bi.read_rows(xlsx(rows), 'students')
        return bi.validate_students(self.school, parsed, create_classes)

    def test_problems_are_reported_per_row(self):
        make_student(self.school, 'TAKEN1', 'Old Student')
        checks, _ = self.check([
            ['Full Name', 'Gender', 'Admission Number', 'Class', 'Parent Email'],
            ['Good Pupil', 'F', '', 'jss 1a', ''],
            ['Bad Pupil', 'X', 'TAKEN1', 'JSS 9Z', 'nope'],
        ])
        self.assertTrue(checks[0].ok)
        errors = ' '.join(checks[1].errors)
        for needle in ('not a gender', 'already used', 'no class called "JSS 9Z"', 'not an email'):
            self.assertIn(needle, errors)

    def test_import_creates_students_with_logins(self):
        checks, new = self.check([
            ['Full Name', 'Class', 'Parent Phone', 'Admission Number'],
            ['Amina Bello', 'JSS 1A', 8031234567, ''],
            ['Musa Bello', 'SS 2B', '', 'MB-1'],
        ], create_classes=True)
        self.assertEqual(new, ['SS 2B'])
        created = bi.import_students(self.school, checks, new)
        amina = Student.objects.get(full_name='Amina Bello')
        self.assertEqual((amina.admission_number, amina.class_assigned, amina.parent_phone),
                         ('STU0001', self.jss1, '08031234567'))
        self.assertEqual(Student.objects.get(admission_number='MB-1').class_assigned.name, 'SS 2B')
        self.assertEqual(created[0]['student_password'], amina.user.term_password)
        self.assertTrue(amina.user.check_password(amina.user.term_password))
        self.assertEqual(amina.parent_term_password, created[0]['parent_password'])

    def test_generated_admission_numbers_never_reuse_one(self):
        make_student(self.school, 'STU0007', 'Old Student')
        self.assertEqual(bi.next_admission_numbers(self.school, 2), ['STU0008', 'STU0009'])

    def test_duplicates_within_the_file(self):
        checks, _ = self.check([['Full Name', 'Admission Number'], ['A One', 'X1'], ['B Two', 'x1']])
        self.assertTrue(checks[0].ok)
        self.assertIn('also on row 2', checks[1].errors[0])

    def test_the_whole_flow_through_the_pages(self):
        r = self._get('/import/students/template.xlsx')
        self.assertIn('Students', openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames)
        r = self._post('/students/import/', {'file': xlsx([['Full Name', 'Class'], ['Amina Bello', 'JSS 1A'],
                                                           ['Nobody', 'JSS 9Z']])})
        self.assertContains(r, 'Import 1 student')
        self.assertFalse(Student.objects.filter(full_name='Amina Bello').exists())     # preview saves nothing
        r = self._post('/students/import/confirm/')
        self.assertEqual(r['Location'], '/students/import/done/')
        amina = Student.objects.get(full_name='Amina Bello')
        self.assertFalse(Student.objects.filter(full_name='Nobody').exists())
        self.assertContains(self._get('/students/import/done/'), amina.admission_number)
        sheet = openpyxl.load_workbook(io.BytesIO(self._get('/import/students/logins.xlsx').content)).active
        self.assertEqual(sheet.cell(row=5, column=1).value, amina.admission_number)
        self.client.logout()
        r = self.client.post('/student/login/', {'admission_number': amina.admission_number,
                                                 'password': amina.user.term_password}, HTTP_HOST=self.host(self.school))
        self.assertEqual(r['Location'], '/dashboard/')

    def test_teachers_cannot_import(self):
        teacher = make_teacher(self.school)
        self.assertEqual(self._get('/students/import/', user=teacher).status_code, 302)
        self.assertEqual(self._get('/import/students/template.xlsx', user=teacher).status_code, 302)


class StaffTests(Base):

    def check(self, rows):
        parsed, _ = bi.read_rows(xlsx(rows), 'staff')
        return bi.validate_staff(self.school, parsed)

    def test_roles_usernames_and_sections(self):
        Campus.objects.create(tenant=self.school, name='Secondary')
        checks = self.check([
            ['Full Name', 'Role', 'Username', 'Section', 'Phone'],
            ['Chidi Okafor', 'Teacher', '', 'secondary', '08031112222'],
            ['Chidi Okafor', 'Principal', '', '', ''],
            ['Aisha Danjuma', 'Cleaner', 'admin1', 'Annex', ''],
        ])
        self.assertEqual((checks[0].data['username'], checks[0].data['section']), ('chidi.okafor', 'Secondary'))
        self.assertEqual((checks[1].data['username'], checks[1].data['role']), ('chidi.okafor2', 'ADMIN'))
        self.assertIn('not a role', ' '.join(checks[2].errors))
        self.assertIn('no section called "Annex"', ' '.join(checks[2].errors))

    def test_import_creates_staff_who_must_change_password(self):
        Campus.objects.create(tenant=self.school, name='Secondary')
        checks = self.check([['Full Name', 'Role', 'Section', 'Phone', 'Job Title'],
                             ['Chidi Okafor', 'Teacher', 'Secondary', '08031112222', 'Maths Teacher']])
        created = bi.import_staff(self.school, checks)
        staff = Staff.objects.get(full_name='Chidi Okafor')
        self.assertTrue(staff.user.must_change_password)
        self.assertTrue(staff.user.check_password(created[0]['temporary_password']))
        self.assertEqual((staff.record.phone, staff.record.job_title), ('08031112222', 'Maths Teacher'))
        self.assertEqual([c.name for c in staff.campuses.all()], ['Secondary'])

    def test_staff_sign_in_then_change_password(self):
        self._post('/staff/import/', {'file': xlsx([['Full Name', 'Role'], ['Chidi Okafor', 'Teacher']])})
        self._post('/staff/import/confirm/')
        temp = self.client.session['import_staff_done']['created'][0]['temporary_password']
        self.client.logout()
        r = self.client.post('/staff/login/', {'username': 'chidi.okafor', 'password': temp},
                             HTTP_HOST=self.host(self.school))
        self.assertEqual(r['Location'], '/change-password/')
