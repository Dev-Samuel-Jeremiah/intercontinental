"""
users/bulk_import.py
────────────────────
Registering many students or staff at once from an Excel (or CSV) file.

    1. template_workbook()  the school downloads a ready-made spreadsheet
    2. read_rows()          the filled-in file is read (any column order; headers matched loosely)
    3. validate_*()         every row is checked and nothing is saved — the admin sees a preview
    4. import_*()           the ready rows are saved in one transaction (all or nothing)
    5. credentials_workbook()  the new logins, to print or share

Accounts get a properly salted PBKDF2 password hash made with fewer rounds, so
hundreds of accounts can be created in one request; Django upgrades each hash
to full strength the first time it is used. Staff must change their temporary
password when they first sign in.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from django.contrib.auth.hashers import PBKDF2PasswordHasher, get_hasher, make_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction

from .models import GENDER_OPTIONS, STATUS_OPTIONS, Class, Staff, Student, User
from .passwords import generate_term_password

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = {'students': 2000, 'staff': 500}
EXTENSIONS = ('.xlsx', '.xlsm', '.xls', '.csv')


class ImportFileError(ValueError):
    """The file as a whole cannot be read (wrong type, no header row, too big)."""


@dataclass
class Column:
    key: str
    header: str
    required: bool = False
    help: str = ''
    example: str = ''
    aliases: tuple = ()
    width: int = 18
    text: bool = False          # keep as text in Excel (phone numbers, admission numbers)


STUDENT_COLUMNS = [
    Column('full_name', 'Full Name', True, "The pupil's full name, surname last or first — as you want it printed.",
           'Amina Bello', ('name', 'student name', 'pupil name', 'names'), 28),
    Column('gender', 'Gender', False, 'Male or Female (M / F also work).', 'Female', ('sex',), 11),
    Column('date_of_birth', 'Date of Birth', False, 'A date cell, or typed as DD/MM/YYYY.', '14/03/2014',
           ('dob', 'birth date', 'date of birth (dd/mm/yyyy)'), 15),
    Column('admission_number', 'Admission Number', False,
           'Leave blank and the system gives the next free number (STU0001, STU0002 …).', 'ISK/2026/001',
           ('admission no', 'admission no.', 'adm no', 'reg no', 'registration number'), 19, True),
    Column('class_name', 'Class', False, 'Exactly as the class is named in the portal (see the drop-down).', 'JSS 1A',
           ('class name', 'class assigned', 'arm'), 14),
    Column('status', 'Status', False, 'Active (default), Suspended, Withdrawn, Graduated.', 'Active', (), 12),
    Column('parent_name', 'Parent Name', False, 'Parent or guardian.', 'Mr Musa Bello',
           ('parent', 'guardian', 'guardian name', "parent's name", 'father name', 'mother name'), 24),
    Column('parent_phone', 'Parent Phone', False, 'Nigerian numbers such as 0803 123 4567 or +234 803 123 4567.',
           '08031234567', ('phone', 'guardian phone', "parent's phone", 'phone number', 'parent phone number'), 16, True),
    Column('parent_email', 'Parent Email', False, 'Siblings with the same parent email are linked in the parent portal.',
           'bello.family@gmail.com', ('email', 'guardian email', "parent's email"), 26),
    Column('address', 'Address', False, 'Home address.', '12 Zoo Road, Kano', ('home address',), 30),
    Column('nin', 'NIN', False, 'National Identification Number — 11 digits (optional).', '12345678901',
           ('nin number',), 14, True),
]

STAFF_ROLES = dict(Staff.ROLE_CHOICES)
STAFF_COLUMNS = [
    Column('full_name', 'Full Name', True, 'As it should appear on ID cards and reports.', 'Chidi Okafor',
           ('name', 'staff name', 'names'), 26),
    Column('role', 'Role', True, 'Teacher, Administrator, Account Officer or Librarian.', 'Teacher',
           ('position', 'staff role', 'designation'), 16),
    Column('username', 'Username', False,
           'Leave blank and one is made from the name (chidi.okafor). Letters, numbers, dot, dash, underscore.',
           'chidi.okafor', ('login', 'user name'), 18, True),
    Column('email', 'Email', False, 'Used for password resets and notices.', 'c.okafor@gmail.com', ('email address',), 26),
    Column('phone', 'Phone', False, 'Nigerian numbers such as 0803 123 4567.', '08031234567',
           ('phone number', 'mobile', 'gsm'), 16, True),
    Column('job_title', 'Job Title', False, 'e.g. Mathematics Teacher, Bursar.', 'Mathematics Teacher',
           ('title', 'job'), 22),
    Column('section', 'Section', False, 'The school section they work in (see the drop-down).', 'Secondary',
           ('campus', 'school section', 'wing'), 18),
    Column('date_joined', 'Date Joined', False, 'DD/MM/YYYY.', '05/09/2022', ('start date', 'employment date'), 14),
    Column('date_of_birth', 'Date of Birth', False, 'DD/MM/YYYY — used for birthday greetings.', '21/07/1988',
           ('dob', 'birth date'), 14),
]

COLUMNS = {'students': STUDENT_COLUMNS, 'staff': STAFF_COLUMNS}


# ─────────────────────────────────────────────────────────────────────────────
# Cleaning single values
# ─────────────────────────────────────────────────────────────────────────────

def _norm_header(text):
    return re.sub(r'[^a-z0-9 ]+', '', str(text or '').replace('*', '').lower()).strip()


def _text(value):
    """Cell value → trimmed text. Whole-number floats lose their '.0' (Excel stores 8031234567 as a float)."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    return ' '.join(str(value).split())


def clean_phone(raw):
    raw = _text(raw)
    if not raw:
        return ''
    digits = re.sub(r'[^\d+]', '', raw)
    if digits.startswith('+234'):
        digits = '0' + digits[4:]
    elif digits.startswith('234') and len(digits) == 13:
        digits = '0' + digits[3:]
    elif len(digits) == 10 and digits[0] in '789':
        digits = '0' + digits                      # Excel dropped the leading 0
    if not re.fullmatch(r'\+?\d{7,15}', digits):
        raise ValueError(f'"{raw}" is not a phone number.')
    if len(digits) == 11 and not digits.startswith('0'):
        raise ValueError(f'"{raw}" looks wrong — Nigerian numbers start with 0 (e.g. 0803…).')
    return digits


DATE_FORMATS = ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d/%m/%y', '%d %b %Y', '%d %B %Y', '%b %d %Y')


def clean_date(raw, *, earliest=1940):
    if raw in (None, ''):
        return None
    if isinstance(raw, datetime):
        d = raw.date()
    elif isinstance(raw, date):
        d = raw
    elif isinstance(raw, (int, float)) and 1000 < raw < 80000:          # an Excel serial number
        d = date(1899, 12, 30) + timedelta(days=int(raw))
    elif re.fullmatch(r'\d{4,5}(\.0)?', _text(raw)) and 1000 < float(_text(raw)) < 80000:
        d = date(1899, 12, 30) + timedelta(days=int(float(_text(raw))))     # a date that arrived as a number
    else:
        text = _text(raw).replace(',', '')
        for fmt in DATE_FORMATS:
            try:
                d = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f'"{text}" is not a date — use DD/MM/YYYY, e.g. 14/03/2014.')
    if d > date.today():
        raise ValueError(f'{d:%d/%m/%Y} is in the future.')
    if d.year < earliest:
        raise ValueError(f'{d:%d/%m/%Y} is too long ago — check the year.')
    return d


def _choice(raw, options, what):
    text = _text(raw)
    if not text:
        return ''
    low = text.lower()
    for value, label in options:
        if low in (value.lower(), label.lower()):
            return value
    if what == 'gender' and low in ('m', 'f'):
        return 'Male' if low == 'm' else 'Female'
    if what == 'role':
        alias = {'admin': 'ADMIN', 'administration': 'ADMIN', 'principal': 'ADMIN', 'head teacher': 'ADMIN',
                 'headteacher': 'ADMIN', 'vice principal': 'ADMIN', 'accountant': 'ACCOUNT', 'bursar': 'ACCOUNT',
                 'account': 'ACCOUNT', 'accounts': 'ACCOUNT', 'library': 'LIB', 'lib': 'LIB',
                 'teachers': 'TEACHER', 'tutor': 'TEACHER', 'class teacher': 'TEACHER'}
        if low in alias:
            return alias[low]
    allowed = ', '.join(label for _v, label in options)
    raise ValueError(f'"{text}" is not a {what} — use one of: {allowed}.')


def quick_hash(raw):
    """A salted PBKDF2 hash with fewer rounds; Django upgrades it on first successful sign-in.
    If the site uses a different password hasher, hash normally with that one."""
    hasher = get_hasher('default')
    if isinstance(hasher, PBKDF2PasswordHasher):
        return hasher.encode(raw, hasher.salt(), iterations=20000)
    return make_password(raw)


# ─────────────────────────────────────────────────────────────────────────────
# Reading the file
# ─────────────────────────────────────────────────────────────────────────────

def _grid(upload):
    """All rows of the first relevant sheet, as lists of raw cell values."""
    name = (upload.name or '').lower()
    if not name.endswith(EXTENSIONS):
        raise ImportFileError('Upload an Excel file (.xlsx or .xls) or a CSV file.')
    if upload.size > MAX_FILE_BYTES:
        raise ImportFileError('The file is larger than 5 MB.')
    data = upload.read()
    try:
        if name.endswith('.csv'):
            text = data.decode('utf-8-sig', errors='replace')
            return list(csv.reader(io.StringIO(text)))
        if name.endswith('.xls'):
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            return rows
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = next((ws for ws in wb.worksheets if ws.title.lower() in ('students', 'staff')), wb.worksheets[0])
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError('This file could not be opened. Save it again from Excel as .xlsx and retry.')


def read_rows(upload, kind):
    """[(spreadsheet row number, {column key: raw value})] plus a list of headers we did not recognise."""
    columns = COLUMNS[kind]
    lookup = {}
    for col in columns:
        for name in (col.header, col.key.replace('_', ' '), *col.aliases):
            lookup[_norm_header(name)] = col.key
    grid = _grid(upload)
    header_at, mapping = None, {}
    for i, row in enumerate(grid[:10]):                      # the header is in the first few rows
        found = {j: lookup[_norm_header(v)] for j, v in enumerate(row) if _norm_header(v) in lookup}
        if 'full_name' in found.values():
            header_at, mapping = i, found
            break
    if header_at is None:
        raise ImportFileError('No header row found. Use the template, or make sure one column is headed "Full Name".')
    missing = [c.header for c in columns if c.required and c.key not in mapping.values()]
    if missing:
        raise ImportFileError(f'The file has no {" or ".join(missing)} column.')
    unknown = [_text(v) for j, v in enumerate(grid[header_at]) if _text(v) and j not in mapping]
    rows = []
    for offset, raw in enumerate(grid[header_at + 1:], start=header_at + 2):
        values = {key: raw[j] if j < len(raw) else None for j, key in mapping.items()}
        if not any(_text(v) for v in values.values()):
            continue                                          # blank line
        rows.append((offset, {k: (v.strftime('%Y-%m-%d') if isinstance(v, (datetime, date)) else _text(v))
                              for k, v in values.items()}))
    if not rows:
        raise ImportFileError('The file has a header row but no names under it.')
    if len(rows) > MAX_ROWS[kind]:
        raise ImportFileError(f'That is {len(rows)} rows — import at most {MAX_ROWS[kind]} at a time '
                              f'(split the file).')
    return rows, unknown


# ─────────────────────────────────────────────────────────────────────────────
# Checking (nothing is saved)
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class RowCheck:
    number: int
    raw: dict
    data: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


def _check(row, key, fn, *args, **kw):
    try:
        row.data[key] = fn(row.raw.get(key, ''), *args, **kw)
    except ValueError as problem:
        row.errors.append(str(problem))
        row.data[key] = None


def validate_students(tenant, rows, create_classes=False):
    classes = {c.name.strip().lower(): c for c in Class.objects.filter(tenant=tenant)}
    taken = {a.lower() for a in Student.objects.filter(tenant=tenant).values_list('admission_number', flat=True)}
    taken |= {u.lower() for u in User.objects.filter(tenant=tenant).values_list('username', flat=True)}
    known = {(n.lower(), d): a for n, d, a in Student.objects.filter(tenant=tenant)
             .values_list('full_name', 'date_of_birth', 'admission_number')}
    seen_adm, new_classes, checks = {}, set(), []
    for number, raw in rows:
        r = RowCheck(number, raw)
        name = _text(raw.get('full_name'))
        if not name:
            r.errors.append('Full Name is empty.')
        elif len(name) > 255:
            r.errors.append('Full Name is too long.')
        r.data['full_name'] = name
        _check(r, 'gender', _choice, GENDER_OPTIONS, 'gender')
        _check(r, 'date_of_birth', clean_date, earliest=1990)
        _check(r, 'status', _choice, STATUS_OPTIONS, 'status')
        r.data['status'] = r.data.get('status') or 'Active'
        adm = _text(raw.get('admission_number'))
        if adm:
            if len(adm) > 30:
                r.errors.append('Admission Number is longer than 30 characters.')
            elif adm.lower() in taken:
                r.errors.append(f'Admission Number {adm} is already used in the portal.')
            elif adm.lower() in seen_adm:
                r.errors.append(f'Admission Number {adm} is also on row {seen_adm[adm.lower()]}.')
            seen_adm.setdefault(adm.lower(), number)
        r.data['admission_number'] = adm
        cls = _text(raw.get('class_name'))
        r.data['class_name'] = cls
        if cls and cls.lower() not in classes:
            if create_classes:
                new_classes.add(cls)
                r.warnings.append(f'Class "{cls}" will be created.')
            else:
                close = ', '.join(sorted(c.name for c in classes.values())[:12]) or 'none yet'
                r.errors.append(f'There is no class called "{cls}". Classes in the portal: {close}.')
        elif not cls:
            r.warnings.append('No class — the student can be placed in a class later.')
        r.data['parent_name'] = _text(raw.get('parent_name'))[:255]
        _check(r, 'parent_phone', clean_phone)
        email = _text(raw.get('parent_email')).lower()
        if email:
            try:
                validate_email(email)
            except ValidationError:
                r.errors.append(f'"{email}" is not an email address.')
        r.data['parent_email'] = email
        r.data['address'] = _text(raw.get('address'))
        nin = re.sub(r'\D', '', _text(raw.get('nin')))
        if nin and len(nin) != 11:
            r.errors.append('NIN must be 11 digits.')
        r.data['nin'] = nin
        if name and (name.lower(), r.data.get('date_of_birth')) in known and r.data.get('date_of_birth'):
            r.warnings.append(f'Looks like a student already registered '
                              f'({known[(name.lower(), r.data["date_of_birth"])]}).')
        if not r.data['parent_phone'] and not email:
            r.warnings.append('No parent phone or email — the family will not receive messages.')
        checks.append(r)
    return checks, sorted(new_classes, key=str.lower)


_USERNAME_OK = re.compile(r'^[A-Za-z0-9._-]{3,150}$')


def _username_from(name, taken):
    parts = [re.sub(r'[^a-z0-9]', '', p) for p in name.lower().split()]
    parts = [p for p in parts if p]
    base = '.'.join([parts[0], parts[-1]]) if len(parts) > 1 else (parts[0] if parts else 'staff')
    candidate, n = base, 1
    while candidate.lower() in taken:
        n += 1
        candidate = f'{base}{n}'
    return candidate


def validate_staff(tenant, rows):
    from saas_admin.models import Campus
    taken = {u.lower() for u in User.objects.filter(tenant=tenant).values_list('username', flat=True)}
    staff_emails = {e.lower() for e in Staff.objects.filter(tenant=tenant).values_list('user__email', flat=True) if e}
    sections = {c.name.strip().lower(): c.name for c in Campus.objects.filter(tenant=tenant, active=True)}
    in_file, checks = {}, []
    for number, raw in rows:
        r = RowCheck(number, raw)
        name = _text(raw.get('full_name'))
        if not name:
            r.errors.append('Full Name is empty.')
        r.data['full_name'] = name[:255]
        _check(r, 'role', _choice, Staff.ROLE_CHOICES, 'role')
        if not r.data.get('role') and not r.errors:
            r.errors.append('Role is empty — Teacher, Administrator, Account Officer or Librarian.')
        username = _text(raw.get('username'))
        if username:
            if not _USERNAME_OK.match(username):
                r.errors.append(f'Username "{username}" can only use letters, numbers, dot, dash and underscore '
                                f'(at least 3).')
            elif username.lower() in taken:
                r.errors.append(f'Username "{username}" is already used in the portal.')
            elif username.lower() in in_file:
                r.errors.append(f'Username "{username}" is also on row {in_file[username.lower()]}.')
        elif name:
            username = _username_from(name, taken | set(in_file))
            r.warnings.append(f'Username will be {username}.')
        if username:
            in_file.setdefault(username.lower(), number)
        r.data['username'] = username
        email = _text(raw.get('email')).lower()
        if email:
            try:
                validate_email(email)
            except ValidationError:
                r.errors.append(f'"{email}" is not an email address.')
            if email in staff_emails:
                r.warnings.append(f'A staff member with {email} already exists.')
        r.data['email'] = email
        _check(r, 'phone', clean_phone)
        r.data['job_title'] = _text(raw.get('job_title'))[:100]
        section = _text(raw.get('section'))
        if section and section.lower() not in sections:
            r.errors.append(f'There is no section called "{section}". Sections: '
                            f'{", ".join(sorted(sections.values())) or "none yet"}.')
        r.data['section'] = sections.get(section.lower(), '') if section else ''
        _check(r, 'date_joined', clean_date, earliest=1960)
        _check(r, 'date_of_birth', clean_date, earliest=1940)
        checks.append(r)
    return checks


# ─────────────────────────────────────────────────────────────────────────────
# Saving
# ─────────────────────────────────────────────────────────────────────────────

def next_admission_numbers(tenant, count, reserved=()):
    """`count` fresh STU#### numbers after the highest one in use (never reusing a number)."""
    used = {a.upper() for a in Student.objects.filter(tenant=tenant).values_list('admission_number', flat=True)}
    used |= {u.upper() for u in User.objects.filter(tenant=tenant).values_list('username', flat=True)}
    used |= {r.upper() for r in reserved}
    top = max([int(m.group(1)) for a in used if (m := re.fullmatch(r'STU(\d+)', a))] or [0])
    out, n = [], top
    while len(out) < count:
        n += 1
        candidate = f'STU{n:04d}'
        if candidate not in used:
            out.append(candidate)
    return out


def _first_last(name):
    first, _, last = name.partition(' ')
    return first[:150], last[:150]


def import_students(tenant, checks, new_classes=()):
    """Save every row that passed. Returns one dict per new student, with their passwords."""
    ok = [c for c in checks if c.ok]
    created = []
    with transaction.atomic():
        classes = {c.name.strip().lower(): c for c in Class.objects.filter(tenant=tenant)}
        for name in new_classes:
            if name.lower() not in classes:
                classes[name.lower()] = Class.objects.create(tenant=tenant, name=name)
        blanks = iter(next_admission_numbers(tenant, sum(1 for c in ok if not c.data['admission_number']),
                                             reserved=[c.data['admission_number'] for c in ok]))
        for c in ok:
            d = c.data
            adm = d['admission_number'] or next(blanks)
            student_pw, parent_pw = generate_term_password(), generate_term_password()
            first, last = _first_last(d['full_name'])
            try:
                user = User.objects.create(username=adm, password=quick_hash(student_pw), first_name=first,
                                           last_name=last, is_student=True, is_staff_user=False, tenant=tenant)
            except IntegrityError:
                raise ImportFileError(f'Row {c.number}: {adm} was registered by someone else meanwhile — '
                                      f'nothing was imported. Upload the file again.')
            user.term_password = student_pw
            user.save(update_fields=['term_password'])
            cls = classes.get(d['class_name'].lower()) if d['class_name'] else None
            Student.objects.create(
                tenant=tenant, user=user, full_name=d['full_name'], gender=d['gender'] or None,
                date_of_birth=d['date_of_birth'], admission_number=adm, class_assigned=cls, status=d['status'],
                parent_name=d['parent_name'] or None, parent_phone=d['parent_phone'] or None,
                parent_email=d['parent_email'] or None, parent_term_password=parent_pw,
                address=d['address'] or None, nin=d['nin'])
            created.append({'admission_number': adm, 'full_name': d['full_name'],
                            'class': cls.name if cls else '', 'student_password': student_pw,
                            'parent_password': parent_pw, 'parent_phone': d['parent_phone'] or ''})
    return created


ROLE_PERMISSIONS = {
    'ADMIN': dict(can_create_staff=True, can_generate_password=True, can_manage_students=True,
                  can_manage_exams=True, can_view_all_exams=True, can_enter_scores=True, can_view_all_results=True),
    'TEACHER': dict(can_manage_students=True, can_enter_scores=True),
}


def import_staff(tenant, checks):
    """Save every row that passed. Returns one dict per new staff member, with their temporary password."""
    from hr.models import StaffRecord
    from saas_admin.models import Campus
    ok = [c for c in checks if c.ok]
    sections = {c.name: c for c in Campus.objects.filter(tenant=tenant, active=True)}
    created = []
    with transaction.atomic():
        for c in ok:
            d = c.data
            temp = generate_term_password(10)
            first, last = _first_last(d['full_name'])
            try:
                user = User.objects.create(username=d['username'], password=quick_hash(temp), email=d['email'],
                                           first_name=first, last_name=last, is_staff=True, is_staff_user=True,
                                           must_change_password=True, tenant=tenant)
            except IntegrityError:
                raise ImportFileError(f'Row {c.number}: username {d["username"]} was taken meanwhile — '
                                      f'nothing was imported. Upload the file again.')
            staff = Staff.objects.create(tenant=tenant, user=user, full_name=d['full_name'], role=d['role'],
                                         **ROLE_PERMISSIONS.get(d['role'], {}))
            if any(d.get(k) for k in ('phone', 'job_title', 'date_joined', 'date_of_birth')):
                StaffRecord.objects.create(tenant=tenant, staff=staff, phone=d['phone'] or '',
                                           job_title=d['job_title'], date_joined=d['date_joined'],
                                           date_of_birth=d['date_of_birth'])
            if d['section'] in sections:
                sections[d['section']].staff.add(staff)
            created.append({'full_name': d['full_name'], 'role': STAFF_ROLES[d['role']], 'username': d['username'],
                            'temporary_password': temp, 'email': d['email'], 'phone': d['phone'] or '',
                            'section': d['section']})
    return created


# ─────────────────────────────────────────────────────────────────────────────
# Spreadsheets out
# ─────────────────────────────────────────────────────────────────────────────

NAVY, YELLOW = '0A1F44', 'F4D21F'


def _book():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    head = dict(font=Font(bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor=NAVY),
                alignment=Alignment(vertical='center', wrap_text=True))
    return wb, head


def template_workbook(kind, tenant):
    """The blank spreadsheet the school fills in, with drop-downs and an instructions sheet."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from saas_admin.models import Campus
    columns = COLUMNS[kind]
    wb, head = _book()
    ws = wb.active
    ws.title = 'Students' if kind == 'students' else 'Staff'
    lists = wb.create_sheet('Lists')
    lists.sheet_state = 'hidden'

    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header + (' *' if col.required else ''))
        for attr, val in head.items():
            setattr(cell, attr, val)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = col.width
        if col.text:
            for r in range(2, MAX_ROWS[kind] + 2):
                ws.cell(row=r, column=i).number_format = '@'
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'B2'

    def dropdown(key, values, list_col):
        if not values:
            return
        for r, v in enumerate(values, start=1):
            lists.cell(row=r, column=list_col, value=v)
        idx = next(i for i, c in enumerate(columns, start=1) if c.key == key)
        letter, list_letter = get_column_letter(idx), get_column_letter(list_col)
        dv = DataValidation(type='list', formula1=f'=Lists!${list_letter}$1:${list_letter}${len(values)}',
                            allow_blank=True, showErrorMessage=False)
        dv.add(f'{letter}2:{letter}{MAX_ROWS[kind] + 1}')
        ws.add_data_validation(dv)

    sections = list(Campus.objects.filter(tenant=tenant, active=True).order_by('name').values_list('name', flat=True))
    if kind == 'students':
        dropdown('gender', [label for _v, label in GENDER_OPTIONS], 1)
        dropdown('class_name', list(Class.objects.filter(tenant=tenant).order_by('name').values_list('name', flat=True)), 2)
        dropdown('status', ['Active', 'Suspended', 'Withdrawn', 'Graduated'], 3)
    else:
        dropdown('role', [label for _v, label in Staff.ROLE_CHOICES], 1)
        dropdown('section', sections, 2)

    info = wb.create_sheet('Instructions', 0)
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 11
    info.column_dimensions['C'].width = 70
    info.column_dimensions['D'].width = 26
    title = info.cell(row=1, column=1, value=f'{tenant.name} — {"student" if kind == "students" else "staff"} import')
    title.font = Font(bold=True, size=14, color=NAVY)
    info.cell(row=2, column=1, value=f'Fill in the "{ws.title}" sheet: one person per row, starting on row 2. '
                                     f'Columns marked * are required. Do not rename the headers.').font = Font(italic=True)
    for i, label in enumerate(('Column', 'Required', 'What to put', 'Example'), start=1):
        cell = info.cell(row=4, column=i, value=label)
        for attr, val in head.items():
            setattr(cell, attr, val)
    for r, col in enumerate(columns, start=5):
        info.cell(row=r, column=1, value=col.header).font = Font(bold=True)
        info.cell(row=r, column=2, value='Yes' if col.required else '')
        info.cell(row=r, column=3, value=col.help).alignment = Alignment(wrap_text=True, vertical='top')
        info.cell(row=r, column=4, value=col.example)
    note = info.cell(row=len(columns) + 6, column=1,
                     value='Upload the file on the import page. You will see every row checked before anything is '
                           'saved.')
    note.font = Font(bold=True, color='8A6D00')
    info.cell(row=len(columns) + 6, column=1).fill = PatternFill('solid', fgColor='FFF6C4')
    wb.active = 1                                            # open on the data sheet
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def credentials_workbook(kind, rows, tenant):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb, head = _book()
    ws = wb.active
    if kind == 'students':
        ws.title = 'New students'
        cols = [('Admission Number', 'admission_number', 18), ('Full Name', 'full_name', 28), ('Class', 'class', 12),
                ('Student Password', 'student_password', 17), ('Parent Password', 'parent_password', 17),
                ('Parent Phone', 'parent_phone', 15)]
        note = ('Students sign in at /student/login/ with their admission number and student password. Parents sign '
                'in at /parent/login/ with the admission number and parent password.')
    else:
        ws.title = 'New staff'
        cols = [('Full Name', 'full_name', 26), ('Role', 'role', 16), ('Username', 'username', 18),
                ('Temporary Password', 'temporary_password', 20), ('Section', 'section', 16), ('Email', 'email', 26),
                ('Phone', 'phone', 15)]
        note = 'Staff sign in at /staff/login/ and must choose their own password the first time.'
    ws.cell(row=1, column=1, value=f'{tenant.name} — new {"student" if kind == "students" else "staff"} logins. '
                                   f'Keep this file private.').font = Font(bold=True, color=NAVY)
    ws.cell(row=2, column=1, value=note).font = Font(italic=True)
    for i, (label, _k, width) in enumerate(cols, start=1):
        cell = ws.cell(row=4, column=i, value=label)
        for attr, val in head.items():
            setattr(cell, attr, val)
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(rows, start=5):
        for i, (_l, key, _w) in enumerate(cols, start=1):
            ws.cell(row=r, column=i, value=row.get(key, '')).number_format = '@'
    ws.freeze_panes = 'A5'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
