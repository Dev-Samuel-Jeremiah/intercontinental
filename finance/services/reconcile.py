"""
finance/services/reconcile.py
─────────────────────────────
Bank reconciliation: does what the bank says match what the school recorded?

The bursar downloads a statement from the bank (CSV or Excel), imports it,
and each line is matched against what the school already has:

  money in   → a family's payment (transfer, POS, cheque or online)
  money out  → a paid expense (the net, after tax withheld), a paid refund,
               or withholding tax paid over to the tax authority

A line matches automatically when its reference turns up on a record, or
when exactly one unmatched record has the same amount within a few days.
Anything doubtful is left for a person — two payments of NGN 50,000 on the
same day are never guessed between. What is left over is the useful part:
money the bank received that nobody recorded, and payments recorded that
never reached the bank.
"""
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import F, Q, Sum
from django.utils import timezone

from finance.models import (BankLine, BankStatement, Expense, Payment, Refund,
                            WhtRemittance)

# How far apart the bank date and the recorded date may be for an amount match.
DATE_WINDOW_DAYS = 3

# Payments that pass through the bank. Cash never appears on a statement.
BANKED_PAYMENT_METHODS = ['BANK_TRANSFER', 'PAYSTACK', 'POS', 'CHEQUE']
BANKED_EXPENSE_METHODS = [Expense.TRANSFER, Expense.CHEQUE]
BANKED_REFUND_METHODS = [Refund.TRANSFER, Refund.CHEQUE]


RECORD_KINDS = ('payment', 'expense', 'refund', 'wht_remittance')


class StatementError(ValueError):
    """The file could not be read as a bank statement."""


# ─────────────────────────────────────────────────────────────────────────────
# Reading the file
# ─────────────────────────────────────────────────────────────────────────────

# Banks name their columns differently; these are the names seen in the wild.
_COLUMNS = {
    'date': ['transaction date', 'trans date', 'txn date', 'date', 'posting date',
             'trans. date', 'value date', 'tran date'],
    'description': ['description', 'narration', 'remarks', 'details', 'transaction details',
                    'particulars', 'narrative', 'transaction description'],
    'reference': ['reference', 'ref', 'reference number', 'ref no', 'ref. no',
                  'transaction reference', 'cheque no', 'chq no', 'session id'],
    'credit': ['credit', 'credits', 'deposit', 'deposits', 'money in', 'cr', 'credit amount',
               'lodgement', 'lodgements'],
    'debit': ['debit', 'debits', 'withdrawal', 'withdrawals', 'money out', 'dr',
              'debit amount'],
    'amount': ['amount', 'transaction amount', 'amt'],
    'balance': ['balance', 'running balance', 'closing balance', 'available balance', 'bal'],
}

_DATE_FORMATS = ['%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d-%b-%Y', '%d-%b-%y', '%d %b %Y',
                 '%d/%m/%y', '%d.%m.%Y', '%d %B %Y', '%b %d, %Y', '%Y/%m/%d']


def _norm(header):
    return re.sub(r'\s+', ' ', str(header or '').strip().lower().replace('_', ' '))


def _find_columns(headers):
    """Map our field names to column positions, or None if it is not a statement."""
    names = [_norm(h) for h in headers]
    found = {}
    for field, options in _COLUMNS.items():
        for option in options:
            if option in names:
                found[field] = names.index(option)
                break
    if 'date' not in found:
        return None
    if 'amount' not in found and not ('credit' in found or 'debit' in found):
        return None
    return found


def parse_amount(raw):
    """'1,250.00' → 1250.00; '(500.00)' and '500.00 DR' → -500.00; blank → None."""
    if raw is None:
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)).quantize(Decimal('0.01'))
    text = str(raw).strip()
    if not text or text in ('-', '—'):
        return None
    negative = text.startswith('(') and text.endswith(')') or text.upper().endswith('DR') \
        or text.startswith('-')
    cleaned = re.sub(r'[^0-9.]', '', text)
    if not cleaned:
        return None
    try:
        value = Decimal(cleaned).quantize(Decimal('0.01'))
    except InvalidOperation:
        return None
    return -value if negative else value


def parse_date(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    # Excel sometimes hands back '2026-09-14 00:00:00'
    text = text.split(' 00:00')[0].split('T')[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _rows_from_file(uploaded):
    """Yield each row of a CSV or Excel file as a list of cells."""
    name = (getattr(uploaded, 'name', '') or '').lower()
    data = uploaded.read()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception:
            raise StatementError('That Excel file could not be opened.')
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
        return
    if name.endswith('.xls'):
        import xlrd
        try:
            sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        except Exception:
            raise StatementError('That Excel file could not be opened.')
        for i in range(sheet.nrows):
            yield sheet.row_values(i)
        return
    for encoding in ('utf-8-sig', 'latin-1'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    yield from csv.reader(io.StringIO(text))


def read_statement(uploaded):
    """
    Turn a bank statement file into a list of dicts:
    {date, description, reference, amount (+in / −out), balance}.

    Banks put a few lines of letterhead above the real header, so the header
    is the first row that looks like one.
    """
    rows = list(_rows_from_file(uploaded))
    columns, start = None, 0
    for i, row in enumerate(rows[:40]):
        columns = _find_columns(row)
        if columns:
            start = i + 1
            break
    if not columns:
        raise StatementError("Couldn't find the column headings. The file needs at least a "
                             "date column and either an amount column or credit and debit "
                             "columns.")

    def cell(row, field):
        index = columns.get(field)
        if index is None or index >= len(row):
            return None
        return row[index]

    lines = []
    for row in rows[start:]:
        if not row or all(c in (None, '') for c in row):
            continue
        when = parse_date(cell(row, 'date'))
        if not when:
            continue            # opening-balance rows, totals, footers
        if 'amount' in columns and 'credit' not in columns:
            amount = parse_amount(cell(row, 'amount'))
        else:
            credit = parse_amount(cell(row, 'credit')) or Decimal('0')
            debit = parse_amount(cell(row, 'debit')) or Decimal('0')
            amount = abs(credit) - abs(debit)
        if not amount:
            continue
        lines.append({
            'date': when,
            'description': str(cell(row, 'description') or '').strip()[:400],
            'reference': str(cell(row, 'reference') or '').strip()[:120],
            'amount': amount,
            'balance': parse_amount(cell(row, 'balance')),
        })
    if not lines:
        raise StatementError('The file has headings but no transactions under them.')
    return lines


@transaction.atomic
def import_statement(tenant, account, uploaded, user=None):
    """Read the file, store every line, and match what can be matched."""
    lines = read_statement(uploaded)

    # A line already imported from an earlier statement is not imported twice.
    existing = set(
        BankLine.objects.filter(tenant=tenant, statement__account=account)
        .values_list('date', 'amount', 'reference', 'description'))
    fresh = [l for l in lines
             if (l['date'], l['amount'], l['reference'], l['description']) not in existing]

    statement = BankStatement.objects.create(
        tenant=tenant, account=account, uploaded_by=user,
        file_name=(getattr(uploaded, 'name', '') or '')[:200],
        period_start=min(l['date'] for l in lines),
        period_end=max(l['date'] for l in lines))
    BankLine.objects.bulk_create([BankLine(tenant=tenant, statement=statement, **l)
                                  for l in fresh])
    matched = auto_match(statement)
    return statement, len(fresh), len(lines) - len(fresh), matched


# ─────────────────────────────────────────────────────────────────────────────
# Matching
# ─────────────────────────────────────────────────────────────────────────────

def _taken(field, tenant_id):
    """Ids already matched to some bank line, so nothing is matched twice.

    Every query here names the school explicitly: the tenant managers return
    all schools' rows when no request is active, and a match must never reach
    into another school's books.
    """
    return set(BankLine.objects.filter(tenant_id=tenant_id, **{f'{field}__isnull': False})
               .values_list(f'{field}_id', flat=True))


def _window(line, days=DATE_WINDOW_DAYS):
    return line.date - timedelta(days=days), line.date + timedelta(days=days)


def _text(line):
    return f'{line.reference} {line.description}'.upper()


def payment_candidates(line, days=DATE_WINDOW_DAYS):
    taken = _taken('payment', line.tenant_id)
    start, end = _window(line, days)
    return (Payment.objects.filter(tenant_id=line.tenant_id, amount=line.amount, method__in=BANKED_PAYMENT_METHODS)
            .exclude(status=Payment.REJECTED).exclude(id__in=taken)
            .filter(created_at__date__gte=start, created_at__date__lte=end)
            .select_related('wallet__student'))


def expense_candidates(line, days=DATE_WINDOW_DAYS):
    taken = _taken('expense', line.tenant_id)
    start, end = _window(line, days)
    # The bank only sees what reached the supplier: the bill less tax withheld.
    return (Expense.objects.filter(tenant_id=line.tenant_id, status=Expense.PAID,
                                   method__in=BANKED_EXPENSE_METHODS)
            .annotate(net=F('amount') - F('wht_amount')).filter(net=-line.amount)
            .exclude(id__in=taken)
            .filter(spent_on__gte=start, spent_on__lte=end))


def refund_candidates(line, days=DATE_WINDOW_DAYS):
    taken = _taken('refund', line.tenant_id)
    start, end = _window(line, days)
    return (Refund.objects.filter(tenant_id=line.tenant_id, amount=-line.amount, status=Refund.PAID,
                                  method__in=BANKED_REFUND_METHODS)
            .exclude(id__in=taken)
            .filter(paid_on__gte=start, paid_on__lte=end)
            .select_related('student'))


def remittance_candidates(line, days=DATE_WINDOW_DAYS):
    taken = _taken('wht_remittance', line.tenant_id)
    start, end = _window(line, days)
    return (WhtRemittance.objects.filter(tenant_id=line.tenant_id, amount=-line.amount,
                                         method__in=BANKED_EXPENSE_METHODS)
            .exclude(id__in=taken)
            .filter(paid_on__gte=start, paid_on__lte=end))


def _by_reference(line):
    """A record whose reference appears on the bank line — the surest match."""
    text = _text(line)
    if line.is_credit:
        taken = _taken('payment', line.tenant_id)
        for payment in (Payment.objects.filter(tenant_id=line.tenant_id, amount=line.amount)
                        .exclude(status=Payment.REJECTED).exclude(id__in=taken)):
            for ref in (payment.reference, payment.paystack_ref):
                if ref and len(ref) >= 6 and ref.upper() in text:
                    return ('payment', payment)
        return None
    taken_e, taken_r = _taken('expense', line.tenant_id), _taken('refund', line.tenant_id)
    taken_w = _taken('wht_remittance', line.tenant_id)
    for expense in (Expense.objects.filter(tenant_id=line.tenant_id, status=Expense.PAID)
                    .annotate(net=F('amount') - F('wht_amount')).filter(net=-line.amount)
                    .exclude(reference='').exclude(id__in=taken_e)):
        if len(expense.reference) >= 4 and expense.reference.upper() in text:
            return ('expense', expense)
    for refund in (Refund.objects.filter(tenant_id=line.tenant_id, amount=-line.amount,
                                         status=Refund.PAID)
                   .exclude(reference='').exclude(id__in=taken_r)):
        if len(refund.reference) >= 4 and refund.reference.upper() in text:
            return ('refund', refund)
    for remittance in (WhtRemittance.objects.filter(tenant_id=line.tenant_id,
                                                    amount=-line.amount)
                       .exclude(reference='').exclude(id__in=taken_w)):
        if len(remittance.reference) >= 4 and remittance.reference.upper() in text:
            return ('wht_remittance', remittance)
    return None


def suggest(line):
    """The single record this line is clearly for, or None when in doubt."""
    hit = _by_reference(line)
    if hit:
        return hit
    if line.is_credit:
        options = list(payment_candidates(line)[:2])
        return ('payment', options[0]) if len(options) == 1 else None
    options = [('expense', e) for e in expense_candidates(line)[:2]] + \
              [('refund', r) for r in refund_candidates(line)[:2]] + \
              [('wht_remittance', w) for w in remittance_candidates(line)[:2]]
    return options[0] if len(options) == 1 else None


def match(line, kind, record, user=None, automatically=False):
    if kind not in RECORD_KINDS:
        raise ValueError('Unknown record type.')
    if record.tenant_id != line.tenant_id:
        raise ValueError('That record belongs to another school.')
    if BankLine.objects.filter(tenant_id=line.tenant_id, **{kind: record}) \
            .exclude(pk=line.pk).exists():
        raise ValueError('That record is already matched to another bank line.')
    line.payment = record if kind == 'payment' else None
    line.expense = record if kind == 'expense' else None
    line.refund = record if kind == 'refund' else None
    line.wht_remittance = record if kind == 'wht_remittance' else None
    line.status = BankLine.MATCHED
    line.matched_automatically = automatically
    line.matched_by = user
    line.matched_at = timezone.now()
    line.save()
    return line


def unmatch(line):
    line.payment = line.expense = line.refund = line.wht_remittance = None
    line.status = BankLine.UNMATCHED
    line.matched_automatically = False
    line.matched_by = None
    line.matched_at = None
    line.save()
    return line


def ignore(line, note='', user=None):
    unmatch(line)
    line.status = BankLine.IGNORED
    line.note = note
    line.matched_by = user
    line.matched_at = timezone.now()
    line.save()
    return line


def auto_match(statement):
    """Match every line that has one clear partner. Returns how many matched."""
    count = 0
    for line in statement.lines.filter(status=BankLine.UNMATCHED).order_by('date', 'id'):
        hit = suggest(line)
        if hit:
            match(line, hit[0], hit[1], automatically=True)
            count += 1
    return count


# ─────────────────────────────────────────────────────────────────────────────
# What is left over
# ─────────────────────────────────────────────────────────────────────────────

def summary(account, start=None, end=None):
    lines = BankLine.objects.filter(tenant_id=account.tenant_id, statement__account=account)
    if start:
        lines = lines.filter(date__gte=start)
    if end:
        lines = lines.filter(date__lte=end)

    def total(qs):
        return qs.aggregate(t=Sum('amount'))['t'] or Decimal('0.00')

    credits = lines.filter(amount__gt=0)
    debits = lines.filter(amount__lt=0)
    unmatched = lines.filter(status=BankLine.UNMATCHED)
    last = lines.exclude(balance__isnull=True).order_by('-date', '-id').first()
    return {
        'lines': lines.count(),
        'matched': lines.filter(status=BankLine.MATCHED).count(),
        'ignored': lines.filter(status=BankLine.IGNORED).count(),
        'unmatched': unmatched.count(),
        'money_in': total(credits),
        'money_out': -total(debits),
        'unmatched_in': total(unmatched.filter(amount__gt=0)),
        'unmatched_out': -total(unmatched.filter(amount__lt=0)),
        'bank_balance': last.balance if last else None,
        'bank_balance_on': last.date if last else None,
    }


def recorded_not_banked(tenant, start, end):
    """
    Payments and spending the school recorded as going through the bank in
    the period, which no statement line accounts for. Each one is either
    still in transit or recorded wrongly.
    """
    payments = (Payment.objects.filter(tenant=tenant, status=Payment.APPROVED,
                                       method__in=BANKED_PAYMENT_METHODS,
                                       created_at__date__gte=start, created_at__date__lte=end)
                .exclude(id__in=_taken('payment', tenant.id)).select_related('wallet__student'))
    expenses = (Expense.objects.filter(tenant=tenant, status=Expense.PAID,
                                       method__in=BANKED_EXPENSE_METHODS,
                                       spent_on__gte=start, spent_on__lte=end)
                .exclude(id__in=_taken('expense', tenant.id)).select_related('category', 'vendor'))
    refunds = (Refund.objects.filter(tenant=tenant, status=Refund.PAID,
                                     method__in=BANKED_REFUND_METHODS,
                                     paid_on__gte=start, paid_on__lte=end)
               .exclude(id__in=_taken('refund', tenant.id)).select_related('student'))
    remittances = (WhtRemittance.objects.filter(tenant=tenant, method__in=BANKED_EXPENSE_METHODS,
                                                paid_on__gte=start, paid_on__lte=end)
                   .exclude(id__in=_taken('wht_remittance', tenant.id)))
    return {'payments': list(payments), 'expenses': list(expenses), 'refunds': list(refunds),
            'remittances': list(remittances)}


def options_for(line, days=30):
    """Records a person might match this line to by hand — same amount, within a month."""
    if line.is_credit:
        return [('payment', p) for p in payment_candidates(line, days)[:15]]
    return ([('expense', e) for e in expense_candidates(line, days)[:15]] +
            [('refund', r) for r in refund_candidates(line, days)[:15]] +
            [('wht_remittance', w) for w in remittance_candidates(line, days)[:15]])
